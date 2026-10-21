"""Data packages: multi-table exports grouped by type / species / source, files and API."""

from __future__ import annotations

import csv
import io
import json
import zipfile

import httpx
import pytest
from openpyxl import load_workbook

from mosaic.api import create_app
from mosaic.datapackage import build_package, species_of, write_package
from mosaic.models import Record
from mosaic.curation import curate

from .test_api import CROSSREF_WORKS, _client  # noqa: F401  (fixtures below reuse the API test setup)
from .conftest import load_fixture


def rec(source, rid, type, **kw) -> Record:
    return Record(source=source, id=rid, type=type, **kw)


RECORDS = [
    rec("ncbi", "1", "article", title="Sotorasib in KRAS G12C NSCLC", date="2021-06-24", authors=["A. Skoulidis"],
        identifiers={"pmid": "34096690", "doi": "10.1056/NEJMoa2103695"},
        attributes={"journal": "N Engl J Med", "mesh": ["Humans", "Lung Neoplasms"], "publication_types": ["Clinical Trial"]}),
    rec("crossref", "10.1056/nejmoa2103695", "article", title="Sotorasib for Lung Cancers with KRAS p.G12C Mutation",
        identifiers={"doi": "10.1056/NEJMoa2103695"}, attributes={"journal": "New England Journal of Medicine", "cited_by": 2500}),
    rec("ncbi", "2", "article", title="AMG 510 in a mouse xenograft model", date="2019",
        identifiers={"pmid": "31666701"}, attributes={"mesh": ["Animals", "Mice", "Humans"]}),
    rec("clinicaltrials", "NCT04303780", "trial", title="CodeBreaK 200", identifiers={"nct": "NCT04303780"},
        attributes={"status": "COMPLETED", "phases": ["PHASE3"], "enrollment": 345}),
    rec("chembl", "CHEMBL4535757", "drug", title="SOTORASIB", identifiers={"chembl": "CHEMBL4535757"},
        attributes={"max_phase": 4, "first_approval": 2021}),
    rec("dandi", "000235", "dataset", title="Mouse visual cortex recordings", identifiers={"dandi": "000235"},
        attributes={"species": ["Mus musculus - House mouse"], "license": "CC-BY-4.0"}),
    rec("uniprot", "P01116", "protein", title="GTPase KRas", identifiers={"uniprot": "P01116"},
        attributes={"organism": "Homo sapiens", "reviewed": True}),
    rec("alliance", "ZFIN:ZDB-GENE-1", "gene", title="kras (zebrafish)", identifiers={"zfin": "ZDB-GENE-1"},
        attributes={"species": "Danio rerio"}),
    rec("ols", "GO:1", "ontology_term", title="Arabidopsis root development", attributes={"organism": "Arabidopsis thaliana"}),
]


def test_species_detection():
    entities = {e.title: e for e in curate(RECORDS).entities}
    assert species_of(entities["AMG 510 in a mouse xenograft model"]) == ["mouse", "human"]  # MeSH check tags
    assert species_of(entities["CodeBreaK 200"]) == ["human"]  # trials are human by definition
    assert species_of(entities["SOTORASIB"]) == ["not_species_specific"]
    assert species_of(entities["Mouse visual cortex recordings"]) == ["mouse"]
    assert species_of(entities["kras (zebrafish)"]) == ["zebrafish"]
    assert species_of(entities["Arabidopsis root development"]) == ["plant"]
    text_only = curate([
        rec("zenodo", "9", "dataset", title="Single-cell atlas of the mouse retina", identifiers={"zenodo": "9"},
            summary="We profiled murine cells after a plant-derived compound and a virus-like particle."),
        rec("openalex", "W9", "article", title="Outcomes in patients with KRAS G12C NSCLC", identifiers={"openalex": "W9"}),
        rec("crossref", "10.1/x", "article", title="A note on methods", identifiers={"doi": "10.1/x"}),
    ]).entities
    by_title = {e.title: species_of(e) for e in text_only}
    assert by_title["Single-cell atlas of the mouse retina"] == ["mouse"]  # not plant / microbe
    assert by_title["Outcomes in patients with KRAS G12C NSCLC"] == ["human"]
    assert by_title["A note on methods"] == ["species_unknown"]


def test_package_by_type_has_one_table_per_kind_with_fitting_columns():
    package = build_package(RECORDS, "type")
    tables = {t.key: t for t in package.tables}
    assert list(tables)[:3] == ["papers", "clinical_trials", "datasets_repositories"]
    assert set(tables) >= {"papers", "clinical_trials", "datasets_repositories", "drugs_compounds", "genes_proteins"}
    papers = tables["papers"]
    assert len(papers.rows) == 2  # NCBI + Crossref copies of the NEJM paper were merged
    assert papers.columns[:3] == ["entity", "title", "date"]
    assert "id_doi" in papers.columns and "id_pmid" in papers.columns and "authors" in papers.columns
    assert "enrollment" not in papers.columns  # trial fields stay in the trial table
    nejm = next(r for r in papers.rows if r["id_doi"] == "10.1056/nejmoa2103695")
    assert nejm["confirmed_by_databases"] == 2 and nejm["quality_grade"] in {"A", "B"}
    assert nejm["sources"].startswith("ncbi:")
    trial = tables["clinical_trials"].rows[0]
    assert trial["id_nct"] == "NCT04303780" and trial["enrollment"] == 345 and trial["phases"] == "PHASE3"


def test_package_by_species_puts_multi_species_papers_in_each():
    package = build_package(RECORDS, "species")
    tables = {t.key: t for t in package.tables}
    assert list(tables)[0] == "human"
    assert {"human", "mouse", "zebrafish", "plant", "not_species_specific"} <= set(tables)
    mouse_titles = {r["title"] for r in tables["mouse"].rows}
    assert mouse_titles == {"AMG 510 in a mouse xenograft model", "Mouse visual cortex recordings"}
    assert "AMG 510 in a mouse xenograft model" in {r["title"] for r in tables["human"].rows}
    columns = tables["mouse"].columns
    assert columns[3] == "type"  # split by species: the computed species column is replaced by type
    assert columns.count("species") == 1  # only the source's own "species" data field (DANDI)


def test_write_package_files_open(tmp_path):
    package = build_package(RECORDS, "type")
    saved = write_package(tmp_path, package, name="KRAS G12C 抑制剂", description="test", query="sotorasib")
    assert saved.folder.name.startswith("KRAS_G12C_抑制剂__") and saved.folder.name.endswith(saved.export_id)
    names = {p.name for p in saved.folder.iterdir()}
    assert {"papers.csv", "clinical_trials.csv", "mosaic_data.xlsx", "README.txt", "manifest.json", "mosaic_data.zip"} <= names

    raw = (saved.folder / "papers.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM so Excel reads UTF-8
    rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    assert len(rows) == 2 and rows[0]["entity"]

    wb = load_workbook(saved.folder / "mosaic_data.xlsx", read_only=True)
    assert wb.sheetnames[0] == "README" and "papers" in wb.sheetnames and "clinical_trials" in wb.sheetnames
    sheet = wb["papers"]
    assert [c.value for c in next(sheet.iter_rows(max_row=1))][:3] == ["entity", "title", "date"]

    with zipfile.ZipFile(saved.folder / "mosaic_data.zip") as archive:
        assert {"papers.csv", "mosaic_data.xlsx", "README.txt", "manifest.json"} <= set(archive.namelist())
    manifest = json.loads((saved.folder / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["group_by"] == "type" and {f["file"] for f in manifest["files"]} >= {"papers.csv"}
    readme = (saved.folder / "README.txt").read_text(encoding="utf-8-sig")
    assert "papers.csv" in readme and "quality_grade" in readme


def test_xlsx_text_never_becomes_a_formula(tmp_path):
    records = [rec("zenodo", "1", "dataset", title="=HYPERLINK(\"http://evil\",\"x\")", identifiers={"zenodo": "1"},
                   attributes={"note": "=1+1", "ctrl": "a\x01b"})]
    saved = write_package(tmp_path, build_package(records, "type"), name="safety")
    wb = load_workbook(saved.folder / "mosaic_data.xlsx")
    ws = wb["datasets_repositories"]
    header = [c.value for c in ws[1]]
    row = {h: c for h, c in zip(header, ws[2])}
    assert row["title"].data_type == "s" and row["note"].data_type == "s"
    assert row["ctrl"].value == "ab"
    csv_text = (saved.folder / "datasets_repositories.csv").read_text(encoding="utf-8-sig")
    assert "'=HYPERLINK" in csv_text  # CSV cells neutralized with a leading apostrophe


# ---------------------------------------------------------------- API
@pytest.fixture
async def pkg_client(offline_settings, make_registry):
    registry = make_registry("crossref")
    app = create_app(offline_settings, registry=registry, agent=None, load_agent_on_startup=False)
    async with _client(app) as c:
        yield c


async def test_search_export_and_download(pkg_client, offline_settings, respx_mock):
    respx_mock.get(CROSSREF_WORKS).mock(return_value=httpx.Response(200, json=load_fixture("crossref", "search_works.json")))
    res = await pkg_client.post("/api/search/export", json={"query": "prime editing", "group_by": "type", "limit_per_source": 2})
    assert res.status_code == 200, res.text
    entry = res.json()
    assert entry["kind"] == "package" and entry["n_rows"] == 2
    files = {f["file"]: f for f in entry["files"]}
    assert set(files) == {"papers.csv"}  # an article and a preprint are both papers

    dl = await pkg_client.get(files["papers.csv"]["url"])
    assert dl.status_code == 200 and dl.headers["content-type"].startswith("text/csv")
    assert "attachment" in dl.headers["content-disposition"]
    assert (await pkg_client.get(entry["xlsx_url"])).headers["content-type"].startswith("application/vnd.openxml")
    assert (await pkg_client.get(entry["zip_url"])).headers["content-type"] == "application/zip"
    preview = (await pkg_client.get(files["papers.csv"]["preview_url"])).json()
    assert preview["columns"][:3] == ["entity", "title", "date"] and len(preview["rows"]) == 2

    listing = (await pkg_client.get("/api/exports")).json()
    assert any(e["export_id"] == entry["export_id"] and e["kind"] == "package" for e in listing["exports"])


@pytest.mark.parametrize("path", [
    "/api/exports/{id}/..%2F..%2Fsecrets.csv",
    "/api/exports/{id}/papers.exe",
    "/api/exports/{id}/../manifest.json",
    "/api/exports/{id}/nope.csv",
    "/api/exports/" + "0" * 32 + "/papers.csv",
    "/api/exports/{id}/mosaic_data.zip/preview",
])
async def test_package_paths_are_validated(pkg_client, offline_settings, path):
    saved = write_package(offline_settings.exports_dir, build_package(RECORDS, "type"), name="t")
    res = await pkg_client.get(path.format(id=saved.export_id))
    assert res.status_code in (400, 404)


def test_comparison_species_have_their_own_buckets():
    cases = {
        "Canis lupus familiaris": "dog", "Sus scrofa": "pig", "Macaca mulatta": "macaque", "Gallus gallus": "chicken",
        "Cavia porcellus guinea pig": "other_animal", "Taeniopygia guttata (zebra finch)": "bird",
        "Pan troglodytes": "primate",
    }
    for organism, bucket in cases.items():
        entity = curate([rec("zenodo", organism, "dataset", title=f"Dataset {organism}", identifiers={"zenodo": organism},
                             attributes={"organism": organism})]).entities[0]
        assert species_of(entity) == [bucket], organism
    package = build_package([rec("chembl", "a1", "activity", title="PK in dogs", attributes={"species": "Canis lupus familiaris"})], "species")
    assert [t.key for t in package.tables] == ["dog"] and package.tables[0].title == "Dog 狗"
