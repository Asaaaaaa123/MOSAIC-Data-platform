"""Full-record enrichment and dataset data files (inventory, safe download, previews, package)."""

from __future__ import annotations

import io
import json
import zipfile

from openpyxl import Workbook, load_workbook

from mosaic.datafiles import ALLOWED_HOSTS, MAX_FILE_BYTES, collect_data_files, inventory
from mosaic.datapackage import build_package, content_metrics, write_package
from mosaic.enrich import enrich_records
from mosaic.http import SourceError
from mosaic.models import Record, ToolResult


class _Conn:
    def __init__(self):
        self.configured = True


class FakeRegistry:
    def __init__(self, handlers):
        self.handlers = handlers
        self.calls = []
        self.connectors = {s: _Conn() for s, _ in handlers}

    async def call(self, source, tool, params, *, session_id=None, timeout=None):
        self.calls.append((source, tool, params))
        handler = self.handlers[(source, tool)]
        out = handler(params)
        if isinstance(out, Exception):
            raise out
        return ToolResult(source=source, tool=tool, records=out)


def hit(source, rid, type, **kw):
    return Record(source=source, id=rid, type=type, **kw)


async def test_enrich_replaces_hits_with_full_records():
    registry = FakeRegistry({
        ("ncbi", "pubmed_fetch"): lambda p: [hit("ncbi", pmid, "article", title=f"Full {pmid}", summary="Full abstract.",
                                                 attributes={"mesh": ["Humans"]}) for pmid in p["pmids"]],
        ("clinicaltrials", "get_study"): lambda p: [hit("clinicaltrials", p["nct"], "trial", title="CodeBreaK 200",
                                                        identifiers={"nct": p["nct"]},
                                                        attributes={"arms": ["sotorasib", "docetaxel"], "enrollment": 345})],
        ("zenodo", "get_record"): lambda p: SourceError("zenodo", "HTTP 503"),
    })
    records = [
        hit("ncbi", "111", "article", title="Hit 111", identifiers={"pmid": "111"}),
        hit("ncbi", "222", "article", title="Hit 222"),
        hit("clinicaltrials", "NCT04303780", "trial", title="CodeBreaK 200 (hit)", attributes={"status": "COMPLETED"}),
        hit("zenodo", "9", "dataset", title="A dataset"),
        hit("bindingdb", "x", "activity", title="no detail tool"),
    ]
    out, stats = await enrich_records(registry, records)
    assert [c[1] for c in registry.calls].count("pubmed_fetch") == 1  # batched
    assert out[0].title == "Full 111" and out[0].summary == "Full abstract." and out[0].identifiers["pmid"] == "111"
    trial = out[2]
    assert trial.attributes["enrollment"] == 345 and trial.attributes["status"] == "COMPLETED"  # hit fields kept
    assert out[3].title == "A dataset"  # failed fetch keeps the hit
    assert stats.enriched == 3 and stats.failed == {"zenodo": 1} and stats.skipped_no_detail_tool == 1
    assert any("3 of 4" in n for n in stats.notes())


async def test_enrich_respects_cap():
    registry = FakeRegistry({("crossref", "get_work"): lambda p: [hit("crossref", p["doi"], "article", title="full")]})
    records = [hit("crossref", f"10.1/{i}", "article") for i in range(10)]
    out, stats = await enrich_records(registry, records, max_records=4)
    assert stats.enriched == 4 and stats.skipped_over_cap == 6
    assert sum(1 for r in out if r.title == "full") == 4


def _xlsx_bytes():
    wb = Workbook()
    ws = wb.active
    ws.append(["gene", "fold_change"])
    ws.append(["KRAS", 2.5])
    ws.append(["TP53", "=1+1"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class FakeHttp:
    def __init__(self, payloads):
        self.payloads = payloads
        self.requested = []

    async def request(self, source, method, url, **kwargs):
        self.requested.append(url)
        data = self.payloads[url]
        if isinstance(data, Exception):
            raise data
        return type("Resp", (), {"content": data})()


DATASET = hit("zenodo", "5022506", "dataset", title="Supplemental data", identifiers={"doi": "10.5281/zenodo.5022506"},
              attributes={"files": [
                  {"name": "counts.csv", "size": 40, "link": "https://zenodo.org/api/records/1/files/counts.csv/content", "checksum": "md5:1"},
                  {"name": "table.xlsx", "size": 5000, "link": "https://zenodo.org/api/records/1/files/table.xlsx/content"},
                  {"name": "raw.fastq.gz", "size": 10_000, "link": "https://zenodo.org/raw"},
                  {"name": "huge.csv", "size": MAX_FILE_BYTES + 1, "link": "https://zenodo.org/huge"},
                  {"name": "secret.csv", "size": 10, "restricted": True, "link": "https://zenodo.org/secret"},
                  {"name": "elsewhere.csv", "size": 10, "link": "https://evil.example/elsewhere.csv"},
                  {"name": "values.json", "size": 30, "link": "https://zenodo.org/values.json"},
              ]})


async def test_collect_data_files_downloads_only_safe_small_open_tables(tmp_path):
    http = FakeHttp({
        "https://zenodo.org/api/records/1/files/counts.csv/content": b"gene,count\nKRAS,10\nTP53,7\n",
        "https://zenodo.org/api/records/1/files/table.xlsx/content": _xlsx_bytes(),
        "https://zenodo.org/values.json": json.dumps([{"a": 1, "b": 2}, {"a": 3, "b": 4}]).encode(),
    })
    result = await collect_data_files(http, [DATASET])
    by_name = {r["file_name"]: r for r in result.rows}
    assert len(by_name) == 7 and "zenodo.org" in ALLOWED_HOSTS
    assert by_name["counts.csv"]["included_in_package"] and by_name["counts.csv"]["preview_sheet"] == "data_1"
    assert by_name["table.xlsx"]["included_in_package"]
    assert "not a tabular" in by_name["raw.fastq.gz"]["note"]
    assert "larger than" in by_name["huge.csv"]["note"]
    assert "restricted" in by_name["secret.csv"]["note"]
    assert "known repository host" in by_name["elsewhere.csv"]["note"]
    assert set(http.requested) == {
        "https://zenodo.org/api/records/1/files/counts.csv/content",
        "https://zenodo.org/api/records/1/files/table.xlsx/content",
        "https://zenodo.org/values.json",
    }
    sheets = {sheet: (header, rows) for sheet, header, rows in result.previews}
    assert sheets["data_1"] == (["gene", "count"], [["KRAS", "10"], ["TP53", "7"]])
    assert sheets["data_2"][0] == ["gene", "fold_change"]

    saved = write_package(tmp_path, build_package([DATASET], "type"), name="with files", data_files=result)
    folder = saved.folder
    assert (folder / "data_files.csv").exists()
    assert (folder / "files" / "zenodo_5022506" / "counts.csv").read_bytes().startswith(b"gene,count")
    with zipfile.ZipFile(folder / "mosaic_data.zip") as archive:
        assert "files/zenodo_5022506/table.xlsx" in archive.namelist() and "data_files.csv" in archive.namelist()
    wb = load_workbook(folder / "mosaic_data.xlsx")
    assert {"data_files", "data_1", "data_2", "data_3"} <= set(wb.sheetnames)
    formula_cell = wb["data_2"]["B3"]
    # downloaded workbooks are read values-only: a formula never reaches the package as a formula
    assert formula_cell.data_type != "f" and formula_cell.value in (None, "=1+1")
    assert wb["data_2"]["B2"].value == 2.5
    readme = (folder / "README.txt").read_text(encoding="utf-8-sig")
    assert "data_files.csv" in readme and "files/" in readme and "% of filled cells are data" in readme


def test_inventory_skips_records_without_files():
    assert inventory([hit("crossref", "10.1/x", "article")]) == []


def test_content_metrics_counts_data_vs_links():
    package = build_package([
        hit("clinicaltrials", "NCT1", "trial", title="T", identifiers={"nct": "NCT1"},
            attributes={"phase": "PHASE3", "enrollment": 300, "results_url": "https://example.org/r"}),
    ], "type")
    metrics = content_metrics(package)
    assert metrics["data_cells"] >= 3 and metrics["identifier_cells"] == 1 and metrics["link_only_cells"] == 1
